import html as _html
import math
import re
import urllib.parse
from datetime import date

import numpy as np
import pandas as pd
import streamlit as st
from fpdf import FPDF
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document
except Exception:
    Document = None


st.set_page_config(page_title="SkillBridge AI", page_icon="🌉", layout="wide")

# ------------------------------- STYLE -------------------------------
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
:root{--text:#EAF0FF;--muted:#A3B0CF;--green:#22E27A;--teal:#2DD4BF;--blue:#3B82F6;--purple:#8B5CF6;--pink:#F472B6;--amber:#F5B544;--line:rgba(255,255,255,.12);}
html, body, .stApp, [class*="css"] { font-family: 'Inter', system-ui, sans-serif; color: var(--text); }
.stApp { background:
radial-gradient(900px 520px at 8% -6%, rgba(45,212,191,.24), transparent 62%),
radial-gradient(800px 600px at 96% 8%, rgba(139,92,246,.26), transparent 62%),
radial-gradient(900px 600px at 50% 112%, rgba(34,226,122,.11), transparent 60%), #060910; }
.block-container { max-width: 1180px; padding-top: 1.8rem; padding-bottom: 4rem; }
#MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; }
/* ---- text visibility: force readable colours on every Streamlit widget ---- */
.stApp [data-testid="stMarkdownContainer"] p, .stApp [data-testid="stMarkdownContainer"] li,
.stApp [data-testid="stMarkdownContainer"] strong, .stApp [data-testid="stMarkdownContainer"] em,
.stApp [data-testid="stMarkdownContainer"] h1, .stApp [data-testid="stMarkdownContainer"] h2,
.stApp [data-testid="stMarkdownContainer"] h3, .stApp [data-testid="stMarkdownContainer"] h4 { color: var(--text) !important; }
.stApp [data-testid="stCaptionContainer"], .stApp [data-testid="stCaptionContainer"] * { color: var(--muted) !important; }
.stApp label, .stApp label *, .stApp [data-testid="stWidgetLabel"], .stApp [data-testid="stWidgetLabel"] * { color: var(--text) !important; }
.stApp a { color: var(--teal) !important; text-decoration: none; }
/* ---- inputs: light background, black text ---- */
.stApp input, .stApp textarea { color: #000 !important; -webkit-text-fill-color: #000 !important; background: #FFFFFF !important; caret-color: #000; }
.stApp input::placeholder, .stApp textarea::placeholder { color: #6B7280 !important; opacity: 1 !important; -webkit-text-fill-color: #6B7280 !important; }
.stApp div[data-baseweb="input"], .stApp div[data-baseweb="base-input"], .stApp div[data-baseweb="textarea"] { background: #FFFFFF !important; border: 1px solid #CBD5E1 !important; border-radius: 10px !important; }
.stApp div[data-baseweb="select"] > div { background: #FFFFFF !important; border: 1px solid #CBD5E1 !important; border-radius: 10px !important; }
.stApp div[data-baseweb="select"] *, .stApp div[data-baseweb="select"] svg { color: #000 !important; -webkit-text-fill-color: #000 !important; fill: #000 !important; }
.stApp div[data-baseweb="select"] input::placeholder { color: #6B7280 !important; -webkit-text-fill-color: #6B7280 !important; }
/* selected skill tags (multiselect) */
span[data-baseweb="tag"] { background: #DCFCE7 !important; border: 1px solid #22C55E !important; }
span[data-baseweb="tag"] *, span[data-baseweb="tag"] svg { color: #064E3B !important; fill: #064E3B !important; -webkit-text-fill-color: #064E3B !important; }
/* dropdown menus */
div[data-baseweb="popover"] > div, div[data-baseweb="popover"] ul, ul[role="listbox"] { background: #FFFFFF !important; border: 1px solid #CBD5E1; border-radius: 12px; }
div[data-baseweb="popover"] li, div[data-baseweb="popover"] li *, li[role="option"], li[role="option"] * { color: #000 !important; -webkit-text-fill-color: #000 !important; background: transparent !important; }
div[data-baseweb="popover"] li:hover, li[role="option"]:hover, li[aria-selected="true"] { background: #DCFCE7 !important; }
[data-testid="stFileUploaderDropzone"] { background: rgba(255,255,255,.05) !important; border: 1px dashed rgba(45,212,191,.55) !important; border-radius: 14px !important; }
[data-testid="stFileUploaderDropzone"], [data-testid="stFileUploaderDropzone"] * { color: var(--text) !important; }
[data-testid="stFileUploaderDropzone"] small { color: var(--muted) !important; }
[data-testid="stFileUploaderDropzone"] button { background: transparent !important; border: 1px solid var(--green) !important; }
[data-testid="stFileUploaderDropzone"] button * { color: var(--green) !important; }
[data-testid="stFileUploaderFile"], [data-testid="stFileUploaderFile"] * { color: var(--text) !important; }
.stApp [data-testid="stAlert"] { background: rgba(255,255,255,.06) !important; border: 1px solid var(--line); border-radius: 12px; }
.stApp [data-testid="stAlert"] * { color: var(--text) !important; }
[data-testid="stExpander"] { background: rgba(255,255,255,.04); border: 1px solid var(--line) !important; border-radius: 14px; }
[data-testid="stExpander"] summary, [data-testid="stExpander"] summary * { color: var(--text) !important; }
hr { border-color: rgba(255,255,255,.12) !important; }
/* ---- headings ---- */
h3 { font-weight: 700 !important; font-size: 1.12rem !important; margin: 1.1rem 0 .5rem !important; padding: 0 !important; letter-spacing: .1px; }
h2 { font-weight: 800 !important; font-size: 1.7rem !important; }
/* ---- brand + hero ---- */
.brand { display: flex; align-items: center; gap: .6rem; font-weight: 700; font-size: 1.05rem; color: var(--text); margin-bottom: 1rem; }
.logo { width: 34px; height: 34px; border-radius: 10px; display: grid; place-items: center; font-weight: 800; font-size: .8rem; color: #04140B; background: linear-gradient(135deg, #2DD4BF, #22E27A); box-shadow: 0 0 18px rgba(34,226,122,.45); }
.hero { font-weight: 800; font-size: 2.9rem; line-height: 1.08; color: var(--text); margin: .2rem 0 .6rem; letter-spacing: -1px; }
.grad { background: linear-gradient(90deg, #2DD4BF, #22E27A 55%, #A3E635); -webkit-background-clip: text; background-clip: text; color: transparent; -webkit-text-fill-color: transparent; }
.sub { color: var(--muted); font-size: 1.02rem; max-width: 640px; line-height: 1.6; margin-bottom: 1.4rem; }
/* ---- glass ---- */
.glass, .st-key-formcard { background: linear-gradient(145deg, rgba(255,255,255,.075), rgba(255,255,255,.025)); border: 1px solid var(--line); border-radius: 20px; backdrop-filter: blur(18px); -webkit-backdrop-filter: blur(18px); box-shadow: 0 24px 60px rgba(0,0,0,.45), inset 0 1px 0 rgba(255,255,255,.09); }
.glass { padding: 1.1rem 1.25rem; margin-bottom: 1rem; color: var(--text); }
.st-key-formcard { padding: 1.4rem 1.6rem 1.2rem; margin-bottom: 1.2rem; }
.ghead { font-weight: 700; font-size: 1.02rem; color: var(--text); margin-bottom: .6rem; }
.gsub { color: var(--muted); font-size: .86rem; margin-bottom: .6rem; line-height: 1.5; }
/* ---- buttons ---- */
.stApp .stButton > button { background: linear-gradient(180deg, #2BEA84, #17C768) !important; border: none !important; border-radius: 12px; font-weight: 700; padding: .7rem 1.2rem; width: 100%; box-shadow: 0 0 22px rgba(34,226,122,.35); }
.stApp .stButton > button, .stApp .stButton > button * { color: #04140B !important; }
.stApp .stButton > button:hover { box-shadow: 0 0 30px rgba(34,226,122,.6); }
.stApp .stDownloadButton > button { background: rgba(34,226,122,.08) !important; border: 1px solid var(--green) !important; border-radius: 12px; font-weight: 700; padding: .7rem 1.2rem; }
.stApp .stDownloadButton > button, .stApp .stDownloadButton > button * { color: var(--green) !important; }
/* ---- tabs ---- */
.stApp div[data-baseweb="tab-list"] { gap: .2rem; border-bottom: 1px solid var(--line); }
.stApp button[data-baseweb="tab"] { background: transparent !important; padding: .55rem .9rem; }
.stApp button[data-baseweb="tab"] p { color: var(--muted) !important; font-weight: 600; font-size: .92rem; }
.stApp button[data-baseweb="tab"][aria-selected="true"] p { color: var(--green) !important; }
.stApp div[data-baseweb="tab-highlight"] { background: var(--green) !important; height: 2px !important; }
.stApp div[data-baseweb="tab-border"] { background: transparent !important; }
/* ---- KPI cards ---- */
.kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 14px; margin: .5rem 0 1.1rem; }
.kpi { background: linear-gradient(145deg, rgba(255,255,255,.075), rgba(255,255,255,.025)); border: 1px solid var(--line); border-radius: 18px; padding: 1rem 1.1rem; box-shadow: 0 16px 40px rgba(0,0,0,.35), inset 0 1px 0 rgba(255,255,255,.08); }
.kpi-l { color: var(--muted); font-size: .8rem; font-weight: 600; letter-spacing: .3px; }
.kpi-v { color: var(--text); font-size: 1.9rem; font-weight: 800; margin: .15rem 0; }
.kpi-d { font-size: .8rem; font-weight: 600; color: var(--green); }
.kpi-d.warn { color: var(--amber); }
.kpi-d.bad { color: var(--pink); }
/* ---- analysis card (gauge + meters) ---- */
.acard { background: linear-gradient(160deg, rgba(255,255,255,.10), rgba(255,255,255,.03)); border: 1px solid rgba(255,255,255,.16); border-radius: 22px; padding: 1.2rem 1.2rem 1.3rem; box-shadow: 0 30px 70px rgba(0,0,0,.5), inset 0 1px 0 rgba(255,255,255,.12); margin-bottom: .8rem; }
.atitle { display: flex; align-items: center; gap: .5rem; font-weight: 700; font-size: 1.05rem; color: var(--text); margin-bottom: 1rem; }
.chipai { font-size: .65rem; font-weight: 800; color: var(--blue); border: 1px solid rgba(59,130,246,.6); border-radius: 5px; padding: .1rem .3rem; }
.ring { width: 210px; height: 210px; border-radius: 50%; margin: 0 auto; display: grid; place-items: center; background: conic-gradient(from 225deg, #22E27A calc(var(--p) * 2.7deg), rgba(255,255,255,.09) 0 270deg, transparent 0); filter: drop-shadow(0 0 14px rgba(34,226,122,.45)); }
.ring-in { width: 168px; height: 168px; border-radius: 50%; background: #0A0F1D; display: grid; place-items: center; align-content: center; }
.ring-num { font-size: 3.6rem; font-weight: 800; color: var(--green); line-height: 1; text-shadow: 0 0 18px rgba(34,226,122,.5); }
.ring-lab { color: var(--muted); font-size: .78rem; font-weight: 600; letter-spacing: 1.5px; margin-top: .2rem; }
.tier { width: fit-content; margin: 1rem auto; padding: .5rem 1.2rem; border-radius: 10px; font-weight: 700; color: var(--green); border: 1px solid rgba(34,226,122,.7); background: rgba(34,226,122,.08); box-shadow: 0 0 20px rgba(34,226,122,.35), inset 0 0 14px rgba(34,226,122,.12); }
.meters { display: grid; grid-template-columns: repeat(4, 1fr); gap: 10px; }
.meter { background: rgba(255,255,255,.05); border: 1px solid var(--line); border-radius: 14px; padding: .7rem .5rem .6rem; text-align: center; }
.mname { color: var(--text); font-weight: 600; font-size: .85rem; margin-bottom: .5rem; }
.segs { display: flex; flex-direction: column-reverse; gap: 5px; margin-bottom: .5rem; }
.seg { height: 7px; border-radius: 4px; background: rgba(255,255,255,.10); }
.seg.on.g { background: #22E27A; box-shadow: 0 0 8px rgba(34,226,122,.55); }
.seg.on.b { background: #3B82F6; box-shadow: 0 0 8px rgba(59,130,246,.55); }
.mval { color: var(--text); font-weight: 700; font-size: .9rem; }
/* ---- chips, bars, tables ---- */
.chip { display: inline-block; padding: .28rem .8rem; margin: .15rem .3rem .15rem 0; border-radius: 999px; font-size: .82rem; font-weight: 600; }
.chip.ok { background: rgba(34,226,122,.14); color: #8FF5BE; border: 1px solid rgba(34,226,122,.5); }
.chip.gap { background: rgba(244,114,182,.14); color: #F9B8D8; border: 1px solid rgba(244,114,182,.5); }
.chip.neutral { background: rgba(255,255,255,.07); color: #D2DBF2; border: 1px solid rgba(255,255,255,.16); }
.chip.info { background: rgba(59,130,246,.15); color: #A9C8FF; border: 1px solid rgba(59,130,246,.5); }
.chip.warn { background: rgba(245,181,68,.14); color: #FBD58B; border: 1px solid rgba(245,181,68,.5); }
.barrow { display: flex; align-items: center; gap: .6rem; min-width: 150px; }
.bar { flex: 1; height: 8px; border-radius: 6px; background: rgba(255,255,255,.09); overflow: hidden; }
.fill { height: 100%; border-radius: 6px; background: linear-gradient(90deg, #2DD4BF, #22E27A); }
.fill.b { background: linear-gradient(90deg, #3B82F6, #8B5CF6); }
.fill.p { background: linear-gradient(90deg, #F472B6, #8B5CF6); }
.barn { color: var(--text); font-weight: 600; font-size: .82rem; min-width: 28px; text-align: right; }
.tbl-wrap { overflow-x: auto; border: 1px solid var(--line); border-radius: 14px; background: rgba(255,255,255,.03); margin: .4rem 0 1rem; }
.tbl { width: 100%; border-collapse: collapse; font-size: .88rem; color: var(--text); }
.tbl th { text-align: left; padding: .65rem .8rem; color: var(--muted); font-weight: 600; font-size: .78rem; letter-spacing: .4px; border-bottom: 1px solid var(--line); white-space: nowrap; }
.tbl td { padding: .65rem .8rem; color: var(--text); border-bottom: 1px solid rgba(255,255,255,.06); vertical-align: middle; }
.tbl tr:last-child td { border-bottom: none; }
.tbl tr:hover td { background: rgba(255,255,255,.03); }
.tbl a { color: var(--teal); font-weight: 600; }
.sec { color: var(--text); font-weight: 700; font-size: 1rem; margin: 1.1rem 0 .4rem; }
.note { color: var(--muted); font-size: .82rem; line-height: 1.5; margin: .3rem 0 .8rem; }
.ck { display: flex; gap: .6rem; align-items: flex-start; padding: .35rem 0; color: var(--text); font-size: .92rem; }
.ck .y { color: var(--green); font-weight: 800; }
.ck .n { color: var(--pink); font-weight: 800; }
.guide { color: var(--text); line-height: 1.7; font-size: .95rem; }
.guide div { margin-bottom: .7rem; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

# ------------------------------- DATA -------------------------------
SKILL_GROUPS = {
    "Programming": ["python", "java", "javascript", "typescript", "c++", "c#", "c programming", "golang", "rust", "php", "ruby", "kotlin", "swift"],
    "Data & AI": ["machine learning", "deep learning", "artificial intelligence", "data science", "data analysis", "data visualization", "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "nlp", "computer vision", "generative ai", "llm", "statistics", "power bi", "tableau", "excel", "sql"],
    "Web": ["html", "css", "react", "angular", "vue", "node.js", "express", "django", "flask", "fastapi", "streamlit", "rest api", "graphql"],
    "Cloud & DevOps": ["aws", "azure", "google cloud", "docker", "kubernetes", "terraform", "jenkins", "ci/cd", "linux", "git", "github", "serverless", "networking"],
    "Databases": ["mysql", "postgresql", "mongodb", "redis", "sqlite", "database design", "data modeling"],
    "Security": ["cybersecurity", "network security", "penetration testing", "ethical hacking", "cryptography", "owasp", "siem"],
    "Core & Professional": ["data structures", "algorithms", "object-oriented programming", "problem solving", "communication", "teamwork", "leadership", "agile", "scrum", "testing", "unit testing", "debugging"],
}
ALL_SKILLS = sorted({s for g in SKILL_GROUPS.values() for s in g}, key=len, reverse=True)

ALIASES = {
    "js": "javascript", "ts": "typescript", "ml": "machine learning", "ai": "artificial intelligence",
    "sklearn": "scikit-learn", "node": "node.js", "nodejs": "node.js", "postgres": "postgresql",
    "powerbi": "power bi", "gcp": "google cloud", "oop": "object-oriented programming",
    "golang": "golang", "cpp": "c++",
}

DISPLAY = {
    "node.js": "Node.js", "ci/cd": "CI/CD", "sql": "SQL", "nlp": "NLP", "llm": "LLM", "aws": "AWS", "html": "HTML",
    "css": "CSS", "rest api": "REST API", "siem": "SIEM", "owasp": "OWASP", "c#": "C#", "c++": "C++", "php": "PHP",
    "power bi": "Power BI", "fastapi": "FastAPI", "mysql": "MySQL", "postgresql": "PostgreSQL", "mongodb": "MongoDB",
    "javascript": "JavaScript", "typescript": "TypeScript", "github": "GitHub", "tensorflow": "TensorFlow",
    "pytorch": "PyTorch", "graphql": "GraphQL", "generative ai": "Generative AI", "scikit-learn": "scikit-learn",
}


def pretty(skill):
    return DISPLAY.get(skill, skill.title())


ROLES = {
    "Data Analyst": {
        "skills": ["sql", "excel", "python", "pandas", "statistics", "power bi", "data visualization", "communication"],
        "projects": ["Build an interactive sales dashboard", "Analyze a public dataset and write up your findings", "Solve a business case study using SQL"],
        "certs": ["Microsoft PL-300: Power BI Data Analyst", "Google Data Analytics Certificate"],
    },
    "Machine Learning Engineer": {
        "skills": ["python", "machine learning", "statistics", "scikit-learn", "pandas", "numpy", "deep learning", "git", "rest api"],
        "projects": ["Train and evaluate a classification model", "Deploy a model behind a REST API", "Build a reproducible end-to-end ML pipeline"],
        "certs": ["Machine Learning Specialization (DeepLearning.AI)", "AWS Certified Machine Learning - Specialty (later stage)"],
    },
    "Python Developer": {
        "skills": ["python", "object-oriented programming", "data structures", "algorithms", "git", "sql", "rest api", "testing", "django"],
        "projects": ["Build a CRUD REST API", "Create a tested automation tool", "Build a Django app with user login"],
        "certs": ["PCAP - Certified Associate Python Programmer", "Python programming specialization"],
    },
    "Full Stack Developer": {
        "skills": ["javascript", "html", "css", "react", "node.js", "rest api", "sql", "git", "testing"],
        "projects": ["Build a full-stack task manager", "Create an e-commerce prototype with login", "Deploy a responsive portfolio with an API"],
        "certs": ["Meta Front-End Developer Certificate", "Full-Stack Web Development specialization"],
    },
    "Cloud Engineer": {
        "skills": ["linux", "aws", "azure", "docker", "kubernetes", "networking", "git", "terraform", "ci/cd"],
        "projects": ["Deploy a containerized web app", "Build a CI/CD pipeline", "Provision cloud infrastructure with Terraform"],
        "certs": ["AWS Certified Cloud Practitioner", "Microsoft Azure Fundamentals (AZ-900)"],
    },
    "Cybersecurity Analyst": {
        "skills": ["network security", "linux", "cybersecurity", "siem", "cryptography", "python", "owasp"],
        "projects": ["Set up a home lab with security monitoring", "Write a threat model for a web app", "Analyze sample logs for suspicious activity"],
        "certs": ["CompTIA Security+", "Google Cybersecurity Certificate"],
    },
    "Software Engineer": {
        "skills": ["python", "java", "data structures", "algorithms", "object-oriented programming", "git", "sql", "testing", "problem solving"],
        "projects": ["Build a well-tested app with a clear architecture", "Solve and document algorithm problems", "Create a REST API with a database"],
        "certs": ["Software design specialization", "Git and GitHub learning path"],
    },
}
CUSTOM = "Paste a job description instead"

COURSES = {
    "python": "Python fundamentals and intermediate programming",
    "sql": "SQL joins, aggregation, and window functions",
    "machine learning": "Supervised learning and model evaluation",
    "pandas": "Data manipulation with pandas",
    "power bi": "Data modeling and dashboards in Power BI",
    "javascript": "Modern JavaScript and async programming",
    "react": "React components, hooks, and state",
    "docker": "Docker images, containers, and Compose",
    "aws": "AWS core services and cloud basics",
    "git": "Git branching and pull requests",
    "statistics": "Probability and hypothesis testing",
    "scikit-learn": "Pipelines and evaluation with scikit-learn",
    "rest api": "Designing, securing, and testing REST APIs",
    "linux": "Linux command line fundamentals",
    "html": "Semantic, accessible HTML",
    "css": "Responsive layouts with modern CSS",
    "data structures": "Arrays, trees, graphs, and complexity",
    "kubernetes": "Kubernetes workloads, services, and deployments",
    "terraform": "Infrastructure as code with Terraform",
    "deep learning": "Neural networks with PyTorch or TensorFlow",
    "excel": "Excel formulas, pivot tables, and charts",
    "data visualization": "Charts and storytelling with data",
    "testing": "Unit and integration testing practices",
    "django": "Building web apps with Django",
    "node.js": "Server-side JavaScript with Node.js",
    "azure": "Azure core services and fundamentals",
    "networking": "TCP/IP, DNS, routing, and firewalls",
    "ci/cd": "Build and deploy pipelines with CI/CD",
    "algorithms": "Sorting, searching, and problem-solving patterns",
    "object-oriented programming": "Classes, inheritance, and design principles",
    "network security": "Defending networks: firewalls, IDS, and VPNs",
    "cybersecurity": "Security fundamentals and threat modeling",
    "siem": "Log analysis and detection with a SIEM",
    "cryptography": "Encryption, hashing, and key management",
    "owasp": "OWASP Top 10 web vulnerabilities",
    "numpy": "Numerical computing with NumPy",
}

DOCS = {
    "python": "https://docs.python.org/3/tutorial/", "docker": "https://docs.docker.com/get-started/",
    "git": "https://git-scm.com/doc", "react": "https://react.dev/learn",
    "pandas": "https://pandas.pydata.org/docs/getting_started/", "scikit-learn": "https://scikit-learn.org/stable/user_guide.html",
    "javascript": "https://developer.mozilla.org/en-US/docs/Learn_web_development", "html": "https://developer.mozilla.org/en-US/docs/Learn_web_development",
    "css": "https://developer.mozilla.org/en-US/docs/Learn_web_development", "kubernetes": "https://kubernetes.io/docs/tutorials/",
    "terraform": "https://developer.hashicorp.com/terraform/tutorials", "aws": "https://aws.amazon.com/training/",
    "django": "https://docs.djangoproject.com/en/stable/intro/tutorial01/", "flask": "https://flask.palletsprojects.com/",
    "fastapi": "https://fastapi.tiangolo.com/tutorial/", "numpy": "https://numpy.org/learn/",
    "streamlit": "https://docs.streamlit.io/", "pytorch": "https://pytorch.org/tutorials/",
    "tensorflow": "https://www.tensorflow.org/tutorials", "node.js": "https://nodejs.org/en/learn",
    "mongodb": "https://learn.mongodb.com/", "postgresql": "https://www.postgresql.org/docs/current/tutorial.html",
    "sql": "https://www.postgresql.org/docs/current/tutorial.html", "owasp": "https://owasp.org/www-project-top-ten/",
    "azure": "https://learn.microsoft.com/training/azure/", "power bi": "https://learn.microsoft.com/power-bi/",
}

CERT_BY_SKILL = {
    "aws": "AWS Certified Cloud Practitioner", "azure": "Microsoft Azure Fundamentals (AZ-900)",
    "power bi": "Microsoft PL-300: Power BI Data Analyst", "kubernetes": "Kubernetes and Cloud Native Associate (KCNA)",
    "terraform": "HashiCorp Terraform Associate", "python": "PCAP - Certified Associate Python Programmer",
    "cybersecurity": "CompTIA Security+", "network security": "CompTIA Security+",
    "machine learning": "Machine Learning Specialization (DeepLearning.AI)", "sql": "Google Data Analytics Certificate",
    "google cloud": "Google Cloud Digital Leader", "git": "Git and GitHub learning path",
}

# Curated, indicative estimates (0-100). Not live market data.
DEMAND = {
    "python": 92, "java": 80, "javascript": 90, "typescript": 85, "c++": 70, "c#": 70, "c programming": 55, "golang": 72,
    "rust": 68, "php": 50, "ruby": 45, "kotlin": 62, "swift": 58, "machine learning": 90, "deep learning": 82,
    "artificial intelligence": 88, "data science": 80, "data analysis": 82, "data visualization": 72, "pandas": 78,
    "numpy": 75, "scikit-learn": 72, "tensorflow": 65, "pytorch": 80, "nlp": 78, "computer vision": 70,
    "generative ai": 92, "llm": 90, "statistics": 78, "power bi": 76, "tableau": 70, "excel": 80, "sql": 94,
    "html": 78, "css": 78, "react": 88, "angular": 65, "vue": 58, "node.js": 80, "express": 66, "django": 66,
    "flask": 62, "fastapi": 72, "streamlit": 55, "rest api": 86, "graphql": 60, "aws": 90, "azure": 82,
    "google cloud": 70, "docker": 88, "kubernetes": 84, "terraform": 74, "jenkins": 56, "ci/cd": 82, "linux": 80,
    "git": 90, "github": 78, "serverless": 62, "networking": 66, "mysql": 70, "postgresql": 78, "mongodb": 68,
    "redis": 62, "sqlite": 45, "database design": 70, "data modeling": 68, "cybersecurity": 84, "network security": 74,
    "penetration testing": 66, "ethical hacking": 62, "cryptography": 58, "owasp": 62, "siem": 64,
    "data structures": 82, "algorithms": 80, "object-oriented programming": 80, "problem solving": 85,
    "communication": 85, "teamwork": 78, "leadership": 76, "agile": 74, "scrum": 64, "testing": 76,
    "unit testing": 70, "debugging": 74,
}
RISING = {"artificial intelligence", "generative ai", "llm", "deep learning", "pytorch", "nlp", "fastapi", "kubernetes",
          "terraform", "rust", "golang", "cybersecurity", "typescript", "ci/cd", "docker", "aws", "azure", "serverless"}
COOLING = {"php", "ruby", "jenkins"}
EFFORT = {  # rough weeks to reach a job-useful level
    "python": 3, "java": 4, "javascript": 3, "sql": 2, "git": 1, "docker": 2, "kubernetes": 4, "aws": 3, "azure": 3,
    "machine learning": 4, "deep learning": 4, "statistics": 3, "pandas": 2, "numpy": 1, "react": 3, "rest api": 2,
    "linux": 2, "html": 1, "css": 1, "power bi": 2, "excel": 1, "data visualization": 2, "terraform": 3, "ci/cd": 2,
    "networking": 3, "testing": 2, "data structures": 3, "algorithms": 3, "django": 3, "node.js": 2,
    "object-oriented programming": 2, "network security": 3, "cybersecurity": 3, "siem": 2, "cryptography": 3,
    "owasp": 2, "scikit-learn": 2, "communication": 2, "problem solving": 3,
}
ROLE_MARKET = {
    "Data Analyst": (80, "Stable", "Steady demand across industries; SQL and visualization are the baseline."),
    "Machine Learning Engineer": (88, "Rising", "Growing with applied AI and LLM adoption; deployment skills stand out."),
    "Python Developer": (84, "Stable", "Broad demand in backend, automation, and data teams."),
    "Full Stack Developer": (82, "Stable", "Consistent demand; React and Node.js dominate listings."),
    "Cloud Engineer": (86, "Rising", "Cloud migration and DevOps practices keep demand high."),
    "Cybersecurity Analyst": (85, "Rising", "Security talent remains in short supply."),
    "Software Engineer": (90, "Stable", "Largest overall pool of openings; fundamentals matter most."),
}

EXPERIENCE_LEVELS = ["Student / Fresher", "0–2 years", "2–5 years", "5+ years"]
EXP_YEARS = {"Student / Fresher": 0, "0–2 years": 1, "2–5 years": 3.5, "5+ years": 6}
LEVEL_TIPS = {
    "Student / Fresher": "Build two or three small projects, host them on GitHub, and add one entry-level certification. Internships and open-source contributions count for a lot.",
    "0–2 years": "Turn your internships or first job into measurable results (for example, time saved or users served) and add one deeper project that shows ownership.",
    "2–5 years": "Show ownership: lead a feature end to end, mention design decisions, and add mentoring or code-review work to your resume.",
    "5+ years": "Highlight system design, leadership, and business impact. Keep one recent hands-on project so your skills read as current.",
}

ACTION = re.compile(r"\b(built|developed|deployed|implemented|designed|created|analy[sz]ed|led|automated|optimi[sz]ed|trained|integrated|managed|engineered)\b")
NICE = re.compile(r"\b(preferred|nice to have|good to have|bonus|a plus|desirable|optional|familiarity)\b", re.I)
SECTION_PATTERNS = {
    "summary": r"summary|objective|profile|about me",
    "education": r"education|academics?|qualifications?",
    "experience": r"(?:work )?experience|internships?|employment",
    "projects": r"(?:academic )?projects?",
    "skills": r"(?:technical )?skills|technologies|tech stack",
    "certifications": r"certifications?|certificates?",
    "achievements": r"achievements?|awards?",
}
KEY_SECTIONS = ["summary", "education", "experience", "projects", "skills"]


# ------------------------------- LOGIC -------------------------------
def esc(x):
    return _html.escape(str(x))


def canon(skill):
    s = skill.strip().lower()
    return ALIASES.get(s, s)


def normalize(text):
    text = (text or "").lower()
    for alias, full in ALIASES.items():
        text = re.sub(r"(?<![\w.+#])" + re.escape(alias) + r"(?![\w.+#])", full, text)
    return text


def skill_pattern(skill):
    return re.compile(r"(?<![a-z0-9+#])" + re.escape(skill) + r"(?![a-z0-9+#])")


def extract_skills(text):
    text = normalize(text)
    return {s for s in ALL_SKILLS if skill_pattern(s).search(text)}


def read_resume(file):
    if file is None:
        return ""
    name = file.name.lower()
    try:
        if name.endswith(".pdf") and PdfReader:
            return "\n".join(p.extract_text() or "" for p in PdfReader(file).pages)
        if name.endswith(".docx") and Document:
            return "\n".join(p.text for p in Document(file).paragraphs)
        if name.endswith(".txt"):
            return file.getvalue().decode("utf-8", errors="ignore")
    except Exception as exc:
        st.warning(f"Couldn't read that file ({exc}). Try a text-based PDF, DOCX or TXT.")
    return ""


def parse_resume(text):
    """Resume parsing: sections, contact presence, years, quantified results, length, clarity score."""
    found = set()
    for line in text.splitlines():
        low = line.strip().lower().rstrip(":")
        if 0 < len(low) <= 40 and len(low.split()) <= 4:
            for key, pat in SECTION_PATTERNS.items():
                if re.match(rf"(?:{pat})\b", low):
                    found.add(key)
    words = len(text.split())
    quantified = len(re.findall(r"\b\d+(?:\.\d+)?\s?(?:%|k\b|x\b|\+)", text))
    years = sorted({int(y) for y in re.findall(r"\b((?:19|20)\d{2})\b", text)})
    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text))
    clarity = 50 * len(found & set(KEY_SECTIONS)) / len(KEY_SECTIONS)
    clarity += 10 if "certifications" in found else 0
    clarity += 15 if quantified >= 2 else 8 if quantified == 1 else 0
    clarity += 15 if has_email else 0
    clarity += 10 if 200 <= words <= 1000 else 0
    return {
        "sections": sorted(found), "words": words, "quantified": quantified,
        "years": years, "has_email": has_email, "clarity": round(min(100.0, clarity), 1),
    }


def skill_strengths(skills, text, claimed):
    """Skill strength (0-100) from how often and how concretely a skill appears in the resume."""
    sentences = re.split(r"[\n.;•]+", normalize(text))
    out = {}
    for s in skills:
        pat = skill_pattern(s)
        mentions = sum(len(pat.findall(x)) for x in sentences)
        applied = any(pat.search(x) and ACTION.search(x) for x in sentences)
        base = {0: 0, 1: 50, 2: 65}.get(mentions, 80)
        if s in claimed:
            base = max(base, 40) + (5 if mentions else 0)
        if applied:
            base += 15
        out[s] = min(100, base)
    return out


def strength_label(v):
    return "Strong" if v >= 75 else "Moderate" if v >= 55 else "Basic"


def analyze_jd(jd):
    """Job description analysis: must-have vs nice-to-have skills and years asked."""
    must, nice = {}, {}
    for sent in re.split(r"[\n.;]+", jd or ""):
        found = extract_skills(sent)
        target = nice if NICE.search(sent) else must
        for s in found:
            target[s] = target.get(s, 0) + 1
    nice = {k: v for k, v in nice.items() if k not in must}
    yrs = [int(y) for y in re.findall(r"(\d+)\s*\+?\s*(?:years|yrs)", jd or "", flags=re.I)]
    return must, nice, (max(yrs) if yrs else None)


def reqs_for_role(role, extra=()):
    skills = ROLES[role]["skills"]
    n = len(skills)
    reqs = {s: round(1.0 - 0.5 * i / max(1, n - 1), 3) for i, s in enumerate(skills)}
    for s in extra:
        reqs.setdefault(s, 0.7)
    return reqs


def reqs_for_jd(jd, extra=()):
    must, nice, _ = analyze_jd(jd)
    reqs = {s: 1.0 for s in must}
    reqs.update({s: 0.5 for s in nice})
    for s in extra:
        reqs.setdefault(s, 0.7)
    return reqs


@st.cache_resource(show_spinner="Loading the SBERT model (first run downloads about 90 MB)...")
def load_sbert():
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer("all-MiniLM-L6-v2")
    except Exception:
        return None


def chunk_resume(text):
    parts = [p.strip(" -•\t") for p in re.split(r"[\n\r]+|(?<=[.!?])\s+", text)]
    return [p for p in parts if len(p) >= 25][:80]


class Semantic:
    """BERT/SBERT semantic matching with a TF-IDF fallback."""

    def __init__(self, text, choice):
        self.text = text
        self.model = None
        self.emb = None
        self.engine = "TF-IDF (keyword similarity)"
        self.chunks = chunk_resume(text) if text.strip() else []
        if choice.startswith("SBERT") and self.chunks:
            model = load_sbert()
            if model is not None:
                try:
                    self.emb = model.encode(self.chunks, normalize_embeddings=True, show_progress_bar=False)
                    self.model = model
                    self.engine = "SBERT (all-MiniLM-L6-v2)"
                except Exception:
                    self.model = None

    def score(self, skills):
        if not self.text.strip() or not skills:
            return {"score": None if not self.text.strip() else 0.0, "per": {}}
        if self.model is not None:
            q = self.model.encode([f"Experience with {s}" for s in skills], normalize_embeddings=True, show_progress_bar=False)
            sims = q @ self.emb.T
            per = {}
            for i, s in enumerate(skills):
                per[s] = float(np.clip((sims[i].max() - 0.15) / 0.45, 0, 1) * 100)
            return {"score": round(float(np.mean(list(per.values()))), 1), "per": per}
        try:
            m = TfidfVectorizer(stop_words="english", ngram_range=(1, 2)).fit_transform([self.text, " ".join(skills)])
            cos = float(cosine_similarity(m[0:1], m[1:2])[0][0])
        except ValueError:
            cos = 0.0
        return {"score": round(min(100.0, cos * 100 / 0.30), 1), "per": {}}


def est_weeks(missing):
    return max(0, round(sum(EFFORT.get(s, 2) for s in missing) * 0.8)) if missing else 0


def combine(reqs, skills, strengths, clarity, semscore):
    """Explainable score: Match, Impact (skill strength), Clarity (resume), Edge (semantic + extras)."""
    total = sum(reqs.values()) or 1.0
    matched = [s for s in reqs if s in skills]
    missing = [s for s in reqs if s not in skills]
    match = 100 * sum(reqs[s] for s in matched) / total if reqs else 0.0
    impact = sum(strengths.get(s, 50) for s in matched) / len(matched) if matched else 0.0
    bonus = min(100.0, 20.0 * len([s for s in skills if s not in reqs]))
    if clarity is None:
        scores = {"match": match, "impact": impact, "clarity": None, "edge": None}
        weights = {"match": 0.70, "impact": 0.30}
    else:
        edge = 0.7 * (semscore or 0.0) + 0.3 * bonus
        scores = {"match": match, "impact": impact, "clarity": clarity, "edge": edge}
        weights = {"match": 0.40, "impact": 0.25, "clarity": 0.15, "edge": 0.20}
    overall = sum(w * scores[k] for k, w in weights.items()) / sum(weights.values())
    return {"matched": matched, "missing": missing, "scores": scores, "weights": weights,
            "overall": round(overall, 1), "semscore": semscore, "bonus": bonus}


def evaluate(name, reqs, ctx):
    skills = list(reqs)
    sem = ctx["sem"].score(skills) if ctx["has_text"] else {"score": None, "per": {}}
    m = combine(reqs, ctx["user_skills"], ctx["strengths"], ctx["clarity"], sem["score"])
    m.update(name=name, reqs=reqs, per=sem["per"], weeks=est_weeks(m["missing"]))
    return m


def prioritize(reqs, missing):
    """Skill prioritization engine: role importance, market demand, and effort."""
    rows = []
    for s in missing:
        w = reqs.get(s, 0.6)
        demand = DEMAND.get(s, 55)
        effort = EFFORT.get(s, 2)
        pr = 50 * w + 35 * demand / 100 + 15 * (1 - (effort - 1) / 3)
        why = []
        if w >= 0.85:
            why.append("core to the role")
        if demand >= 80:
            why.append("high demand")
        if effort <= 1.5:
            why.append("quick win")
        rows.append({"skill": s, "priority": round(pr, 1), "weight": w, "demand": demand, "effort": effort,
                     "why": ", ".join(why) or "supports the role"})
    return sorted(rows, key=lambda r: -r["priority"])


def make_roadmap(priorities):
    steps = []
    for i, p in enumerate(priorities[:9]):
        phase, weeks, milestone = (
            ("Foundations", "Weeks 1–2", "Finish one beginner course or tutorial") if i < 3 else
            ("Practice", "Weeks 3–4", "Build a small hands-on project") if i < 6 else
            ("Portfolio", "Weeks 5–6", "Publish on GitHub and add it to your resume")
        )
        s = p["skill"]
        steps.append({"when": weeks, "focus": phase, "skill": s, "learn": COURSES.get(s, f"Fundamentals and hands-on use of {pretty(s)}"), "milestone": milestone})
    return steps


def learn_links(skill):
    q = urllib.parse.quote_plus(pretty(skill))
    return {
        "docs": DOCS.get(skill),
        "course": f"https://www.coursera.org/search?query={q}",
        "video": f"https://www.youtube.com/results?search_query={q}+tutorial+for+beginners",
    }


def tier(score):
    if score >= 90:
        return "S-Tier", "Excellent Match"
    if score >= 75:
        return "A-Tier", "Strong Match"
    if score >= 60:
        return "B-Tier", "Good Foundation"
    if score >= 40:
        return "C-Tier", "Getting There"
    return "D-Tier", "Early Stage"


def readiness_assessment(m, strengths, parsed, priorities, experience, jd_years):
    reqs = m["reqs"]
    core_missing = [s for s in m["missing"] if reqs[s] >= 0.85]
    avg_strength = (sum(strengths.get(s, 0) for s in m["matched"]) / len(m["matched"])) if m["matched"] else 0
    if m["overall"] >= 80 and not core_missing:
        level = "Job-ready"
    elif m["overall"] >= 65:
        level = "Almost ready"
    elif m["overall"] >= 45:
        level = "Developing"
    else:
        level = "Early stage"
    checks = [
        ("Core skills for the role are covered", not core_missing),
        ("Matched skills show real use (average strength 55 or more)", avg_strength >= 55),
    ]
    if parsed:
        checks += [
            ("Resume has a clear structure (clarity 70 or more)", parsed["clarity"] >= 70),
            ("Resume includes measurable results (numbers or percentages)", parsed["quantified"] >= 2),
            ("Projects section present", "projects" in parsed["sections"]),
            ("Certifications listed", "certifications" in parsed["sections"]),
        ]
    note = None
    if jd_years and EXP_YEARS[experience] + 1 < jd_years:
        note = f"The job description asks for about {jd_years}+ years of experience. Your selected level is lower, so lean on projects and results."
    return {"level": level, "checks": checks, "weeks": m["weeks"], "core_missing": core_missing, "note": note}


def build_recs(m, priorities, role_key, experience, parsed, strengths):
    gaps = [p["skill"] for p in priorities[:6]]
    courses = [(s, COURSES.get(s, f"Fundamentals and hands-on use of {pretty(s)}")) for s in gaps]
    if role_key in ROLES:
        projects = list(ROLES[role_key]["projects"])
        certs = list(ROLES[role_key]["certs"])
    else:
        projects = []
        certs = list(dict.fromkeys(CERT_BY_SKILL[s] for s in m["reqs"] if s in CERT_BY_SKILL))[:3]
    for s in gaps[:3]:
        projects.append(f"Build a small project that shows off {pretty(s)}")
    personal = [LEVEL_TIPS[experience]]
    if gaps:
        personal.append(f"Start with {pretty(gaps[0])}: it ranks highest on role importance, market demand, and effort.")
    weak = sorted((s for s in m["matched"] if strengths.get(s, 0) < 55), key=lambda s: strengths.get(s, 0))[:3]
    if weak:
        personal.append("Deepen skills you list but barely show: " + ", ".join(pretty(s) for s in weak) + ". Use each in a project and describe what you built.")
    if parsed:
        if parsed["quantified"] < 2:
            personal.append("Add measurable results to your resume (for example, 'cut report time by 40%').")
        missing_sections = [s.title() for s in KEY_SECTIONS if s not in parsed["sections"]]
        if missing_sections:
            personal.append("Add clear headings for: " + ", ".join(missing_sections) + ".")
    return {"courses": courses, "projects": projects, "certs": certs, "personal": personal}


def guidance_text(d):
    t = d["target"]
    first = d["name"].split()[0] if d["name"] != "Candidate" else "You"
    tr, label = tier(t["overall"])
    strong = sorted(t["matched"], key=lambda s: -d["strengths"].get(s, 0))[:3]
    paras = [
        f"{first}, your profile scores {t['overall']:.0f} out of 100 for {d['role']} ({tr} / {label}). "
        + (f"Your strongest relevant skills are {', '.join(pretty(s) for s in strong)}." if strong else "No required skills were matched yet, so the first goal is building a base.")
    ]
    if d["priorities"]:
        top = [pretty(p["skill"]) for p in d["priorities"][:3]]
        paras.append(
            f"The best skills to learn next are {', '.join(top)}. They are ranked by how central they are to the role, how much employers ask for them, and how quickly you can learn them. "
            f"A rough estimate for closing your gaps is {t['weeks']} weeks at 8 to 10 hours a week."
        )
    else:
        paras.append("You cover every skill we checked. Focus on portfolio depth, interview practice, and tailoring your resume.")
    if d["parsed"]:
        p = d["parsed"]
        bits = []
        if p["quantified"] < 2:
            bits.append("add measurable results")
        miss = [s for s in KEY_SECTIONS if s not in p["sections"]]
        if miss:
            bits.append("add headings for " + ", ".join(miss))
        paras.append("Your resume clarity is " + f"{p['clarity']:.0f}/100" + (". To improve it, " + " and ".join(bits) + "." if bits else ". The structure looks solid."))
    paras.append(LEVEL_TIPS[d["experience"]])
    others = [c for c in d["comparison"] if c["name"] != t["name"]]
    if others:
        best = max(others, key=lambda c: c["overall"])
        paras.append(f"Another good fit is {best['name']} at {best['overall']:.0f}. It could work as a stepping stone or backup path while you build the skills for {d['role']}.")
    paras.append("This guidance is generated locally from your results using fixed rules. It is a starting point, not a hiring prediction.")
    return paras


# ------------------------------- UI HELPERS -------------------------------
def chips(items, kind):
    return "".join(f'<span class="chip {kind}">{esc(pretty(s))}</span>' for s in items)


def table(headers, rows):
    th = "".join(f"<th>{esc(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
    return f'<div class="tbl-wrap"><table class="tbl"><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table></div>'


def bar(value, kind=""):
    v = max(0, min(100, float(value)))
    return f'<div class="barrow"><div class="bar"><div class="fill {kind}" style="width:{v:.0f}%"></div></div><span class="barn">{v:.0f}</span></div>'


def link(url, text):
    return f'<a href="{esc(url)}" target="_blank" rel="noopener">{esc(text)}</a>'


def show(html_str):
    st.markdown(html_str, unsafe_allow_html=True)


def analysis_card(overall, scores):
    tr, label = tier(overall)
    cols = [("Match", "match", "g"), ("Impact", "impact", "b"), ("Clarity", "clarity", "g"), ("Edge", "edge", "b")]
    meters = ""
    for title, key, color in cols:
        v = scores.get(key)
        on = 0 if v is None else round(v / 100 * 7)
        segs = "".join(f'<div class="seg{" on " + color if i < on else ""}"></div>' for i in range(7))
        meters += f'<div class="meter"><div class="mname">{title}</div><div class="segs">{segs}</div><div class="mval">{"—" if v is None else f"{v:.0f}%"}</div></div>'
    return (
        '<div class="acard"><div class="atitle"><span class="chipai">AI</span>Analysis Result</div>'
        f'<div class="ring" style="--p:{max(0, min(100, overall)):.0f}"><div class="ring-in"><div class="ring-num">{overall:.0f}</div><div class="ring-lab">AI SCORE</div></div></div>'
        f'<div class="tier">{tr} / {label}</div><div class="meters">{meters}</div></div>'
    )


def safe(s):
    s = str(s).replace("–", "-").replace("—", "-").replace("’", "'").replace("“", '"').replace("”", '"').replace("•", "-")
    return s.encode("latin-1", "replace").decode("latin-1")


def build_pdf(d):
    t = d["target"]
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def heading(text):
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, safe(text), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=10)

    def lines(items):
        for item in items or ["None found in the text provided."]:
            pdf.multi_cell(0, 6, "- " + safe(item), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 12, "SkillBridge AI - Career Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 7, f"Generated: {date.today().isoformat()}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 12)
    for line in [f"Name: {d['name']}", f"Target role: {d['role']}", f"Experience: {d['experience']}",
                 f"Overall score: {t['overall']:.0f}/100 ({' / '.join(tier(t['overall']))})",
                 f"Matching engine: {d['engine']}"]:
        pdf.cell(0, 8, safe(line)[:95], new_x="LMARGIN", new_y="NEXT")
    sc = t["scores"]
    heading("Score breakdown")
    lines([f"{k.title()}: {'n/a' if v is None else f'{v:.0f}%'} (weight {t['weights'].get(k, 0) * 100:.0f}%)" for k, v in sc.items()])
    heading("Career readiness")
    lines([f"Level: {d['readiness']['level']}", f"Rough time to close gaps: {t['weeks']} weeks"] + [("[x] " if ok else "[ ] ") + label for label, ok in d["readiness"]["checks"]])
    heading("Matched skills")
    lines([f"{pretty(s)} ({strength_label(d['strengths'].get(s, 0))}, {d['strengths'].get(s, 0):.0f})" for s in t["matched"]])
    heading("Skill gaps in priority order")
    lines([f"{pretty(p['skill'])} - priority {p['priority']:.0f} ({p['why']})" for p in d["priorities"]])
    heading("Six-week roadmap")
    lines([f"{r['when']} | {r['focus']} | {pretty(r['skill'])}: {r['learn']}" for r in d["roadmap"]])
    heading("Course recommendations")
    lines([f"{pretty(s)}: {c}" for s, c in d["recs"]["courses"]])
    heading("Project recommendations")
    lines(d["recs"]["projects"])
    heading("Certification recommendations")
    lines(d["recs"]["certs"])
    heading("Personalized recommendations")
    lines(d["recs"]["personal"])
    heading("Job comparison")
    lines([f"{c['name']}: overall {c['overall']:.0f}, skill match {c['scores']['match']:.0f}%, about {c['weeks']} weeks to close gaps" for c in sorted(d["comparison"], key=lambda c: -c["overall"])])
    heading("Career guidance")
    lines(d["guidance"])
    pdf.ln(2)
    pdf.set_font("Helvetica", "I", 8)
    pdf.multi_cell(0, 5, safe("Indicative report based on the information you provided. It is not a hiring prediction. Market data is a curated estimate."), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ------------------------------- HEADER -------------------------------
show('<div class="brand"><div class="logo">AI</div>SkillBridge AI</div>')
show('<div class="hero"><span class="grad">See how your skills</span><br>match the job you want.</div>')
show('<div class="sub">Upload your resume, pick your skills and a target role. You get an explainable match score, '
     'prioritized skill gaps, learning links, a six-week roadmap, job comparisons, and career guidance.</div>')

# ------------------------------- INPUTS -------------------------------
with st.container(key="formcard"):
    st.markdown("### About you")
    c1, c2 = st.columns(2)
    name = c1.text_input("Full name", placeholder="Alex Morgan")
    experience = c2.selectbox("Experience level", EXPERIENCE_LEVELS)

    st.markdown("### Resume")
    uploaded = st.file_uploader("Upload your resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])

    st.markdown("### Skills you know")
    picked = st.multiselect(
        "Pick from the list",
        options=sorted({pretty(s) for s in ALL_SKILLS}),
        placeholder="Search and select skills",
    )
    other_skills = st.text_input("Other skills (comma separated)", placeholder="Figma, Jira, Spark")

    st.markdown("### Target role")
    c3, c4 = st.columns(2)
    role = c3.selectbox("Role you're aiming for", list(ROLES.keys()) + [CUSTOM])
    engine_choice = c4.selectbox("Matching engine", ["SBERT semantic matching (recommended)", "TF-IDF (fast, no download)"])
    jd = st.text_area(
        "Job description (optional)",
        height=140,
        placeholder="Paste a job description to analyze it. If you chose 'Paste a job description instead' above, it becomes your target.",
    )
    extra_required = st.text_input("Extra skills the job needs (optional, comma separated)", placeholder="Docker, FastAPI")
    more_jds = st.text_area(
        "Compare more jobs (optional)",
        height=110,
        placeholder="Paste extra job descriptions here. Put a line with --- between them.",
    )
    submitted = st.button("Analyze my profile")

# ------------------------------- ANALYSIS -------------------------------
if submitted:
    resume_text = read_resume(uploaded)
    if uploaded is not None and not resume_text.strip():
        st.warning("No text could be read from that file (it may be a scanned image). Using your selected skills instead.")
    manual = {canon(s) for s in other_skills.split(",") if s.strip()}
    claimed = {canon(s) for s in picked} | manual
    user_skills = extract_skills(resume_text) | claimed

    if not resume_text.strip() and not user_skills:
        st.error("Add at least one thing to analyze: upload a resume or pick some skills.")
    elif role == CUSTOM and not jd.strip() and not extra_required.strip():
        st.error("Paste a job description or list the skills the job needs.")
    else:
        extra = sorted({canon(s) for s in extra_required.split(",") if s.strip()})
        has_text = bool(resume_text.strip())
        parsed = parse_resume(resume_text) if has_text else None
        with st.spinner("Analyzing your profile..."):
            ctx = {
                "user_skills": user_skills,
                "strengths": skill_strengths(user_skills, resume_text, claimed),
                "clarity": parsed["clarity"] if parsed else None,
                "has_text": has_text,
                "sem": Semantic(resume_text, engine_choice),
            }
            role_name = role if role != CUSTOM else "Custom job description"
            target_reqs = reqs_for_jd(jd, extra) if role == CUSTOM else reqs_for_role(role, extra)
            target = evaluate(role_name, target_reqs, ctx)

            comparison = []
            for r in ROLES:
                comparison.append(target if r == role else evaluate(r, reqs_for_role(r), ctx))
            if role == CUSTOM:
                comparison.append(target)
            elif jd.strip() and reqs_for_jd(jd):
                comparison.append(evaluate("Pasted job description", reqs_for_jd(jd), ctx))
            for i, chunk in enumerate([c for c in re.split(r"(?m)^\s*-{3,}\s*$", more_jds or "") if c.strip()][:4], start=1):
                r_ = reqs_for_jd(chunk)
                if r_:
                    comparison.append(evaluate(f"Extra job {i}", r_, ctx))

            priorities = prioritize(target["reqs"], target["missing"])
            jd_must, jd_nice, jd_years = analyze_jd(jd)
            recs = build_recs(target, priorities, role, experience, parsed, ctx["strengths"])
            d_new = {
                "name": name.strip() or "Candidate",
                "role": role_name, "role_key": role, "experience": experience,
                "user_skills": user_skills, "strengths": ctx["strengths"], "parsed": parsed,
                "engine": ctx["sem"].engine if has_text else "Skill coverage only (no resume uploaded)",
                "target": target, "comparison": comparison, "priorities": priorities,
                "roadmap": make_roadmap(priorities),
                "jd": {"text": jd.strip(), "must": jd_must, "nice": jd_nice, "years": jd_years},
                "recs": recs, "has_text": has_text,
                "readiness": readiness_assessment(target, ctx["strengths"], parsed, priorities, experience, jd_years),
            }
            d_new["guidance"] = guidance_text(d_new)
        st.session_state["result"] = d_new
        st.session_state["show_guidance"] = False

# ------------------------------- RESULTS -------------------------------
d = st.session_state.get("result")
if d:
    t = d["target"]
    tr, label = tier(t["overall"])
    first = d["name"].split()[0] if d["name"] != "Candidate" else ""
    heading = f"{first}'s results" if first else "Your results"
    st.markdown(f"## {heading} for {d['role']}")

    rd = d["readiness"]["level"]
    kclass = "" if rd in ("Job-ready", "Almost ready") else "warn" if rd == "Developing" else "bad"
    show(
        '<div class="kpis">'
        f'<div class="kpi"><div class="kpi-l">Overall score</div><div class="kpi-v">{t["overall"]:.0f}%</div><div class="kpi-d">{tr} · {label}</div></div>'
        f'<div class="kpi"><div class="kpi-l">Skill match</div><div class="kpi-v">{t["scores"]["match"]:.0f}%</div><div class="kpi-d">{len(t["matched"])} of {len(t["reqs"])} skills</div></div>'
        f'<div class="kpi"><div class="kpi-l">Skills to build</div><div class="kpi-v">{len(t["missing"])}</div><div class="kpi-d warn">about {t["weeks"]} weeks</div></div>'
        f'<div class="kpi"><div class="kpi-l">Readiness</div><div class="kpi-v" style="font-size:1.4rem">{rd}</div><div class="kpi-d {kclass}">for {esc(d["role"])}</div></div>'
        f'<div class="kpi"><div class="kpi-l">Matching engine</div><div class="kpi-v" style="font-size:1.1rem;line-height:1.6">{esc(d["engine"].split(" (")[0])}</div><div class="kpi-d">{"semantic" if d["engine"].startswith("SBERT") else "keyword based"}</div></div>'
        "</div>"
    )

    left, right = st.columns([5, 7], gap="large")
    with left:
        show(analysis_card(t["overall"], t["scores"]))
        if st.button("Generate career guidance ✨"):
            st.session_state["show_guidance"] = True
    with right:
        top = d["priorities"][:5]
        rows = [[f"{i}", esc(pretty(p["skill"])), bar(p["priority"], "b"), esc(p["why"])] for i, p in enumerate(top, 1)]
        show('<div class="glass"><div class="ghead">Top skills to learn next</div>'
             + (table(["#", "Skill", "Priority", "Why"], rows) if rows else '<div class="gsub">No gaps found for this role.</div>') + "</div>")
        best = sorted(d["comparison"], key=lambda c: -c["overall"])[:3]
        show('<div class="glass"><div class="ghead">Best-fitting jobs for you</div>'
             + table(["Job", "Overall", "Skill match"], [[esc(c["name"]), bar(c["overall"]), f'{c["scores"]["match"]:.0f}%'] for c in best]) + "</div>")

    if st.session_state.get("show_guidance"):
        paras = "".join(f"<div>{esc(p)}</div>" for p in d["guidance"])
        show(f'<div class="glass"><div class="ghead">AI career guidance</div><div class="guide">{paras}</div></div>')

    tabs = st.tabs(["Skills", "Explain score", "Roadmap & learning", "Recommendations", "Jobs",
                    "Market trends", "Switch simulator", "Readiness", "Resume & JD", "Download"])

    # ---------- Skills: match, strength, missing, priority ----------
    with tabs[0]:
        show('<div class="sec">Skills you have that this role needs</div>')
        show(chips(t["matched"], "ok") if t["matched"] else '<div class="note">None yet. Add skills or upload a resume.</div>')
        show('<div class="sec">Missing skills</div>')
        show(chips(t["missing"], "gap") if t["missing"] else '<div class="note">You cover every skill we looked for.</div>')
        related = [s for s in t["missing"] if t["per"].get(s, 0) >= 55]
        if related:
            show('<div class="sec">Related experience found by semantic matching</div>')
            show(chips(related, "info"))
            show('<div class="note">Your resume text is close in meaning to these missing skills. Name them explicitly if you have used them.</div>')
        other = sorted(d["user_skills"] - set(t["reqs"]))
        if other:
            show('<div class="sec">Your other skills</div>')
            show(chips(other, "neutral"))

        show('<div class="sec">Skill strength</div>')
        ordered = sorted(d["user_skills"], key=lambda s: (s not in t["reqs"], -d["strengths"].get(s, 0)))
        srows = [[esc(pretty(s)), "Needed" if s in t["reqs"] else "Extra", bar(d["strengths"].get(s, 0), "" if s in t["reqs"] else "b"),
                  strength_label(d["strengths"].get(s, 0))] for s in ordered]
        show(table(["Skill", "Role fit", "Strength", "Level"], srows))
        show('<div class="note">Strength rises with how often a skill appears in your resume, whether it is used in a concrete action (built, deployed, analyzed), and whether you selected it. Skills you only selected start at a basic level.</div>')

        show('<div class="sec">Skill prioritization engine</div>')
        prow = [[str(i), esc(pretty(p["skill"])), bar(p["priority"], "b"), f'{p["demand"]}', f'{p["effort"]} wk', esc(p["why"])]
                for i, p in enumerate(d["priorities"], 1)]
        show(table(["#", "Skill", "Priority", "Demand", "Effort", "Why"], prow) if prow else '<div class="note">Nothing to prioritize.</div>')
        show('<div class="note">Priority = 50% role importance + 35% market demand + 15% ease of learning. Demand and effort are curated estimates.</div>')

    # ---------- Explain score ----------
    with tabs[1]:
        w = t["weights"]
        desc = {"match": "Weighted share of required skills you have", "impact": "Average strength of your matched skills",
                "clarity": "Resume structure, contact info, and measurable results", "edge": "Semantic fit to the role plus extra skills"}
        erows = []
        for k, v in t["scores"].items():
            if k in w:
                erows.append([k.title(), esc(desc[k]), "—" if v is None else f"{v:.0f}", f"{w[k] * 100:.0f}%", f"{w[k] * (v or 0):.1f}"])
        erows.append(["<b>Overall</b>", "", "", "100%", f"<b>{t['overall']:.1f}</b>"])
        show(table(["Component", "What it measures", "Score", "Weight", "Points"], erows))
        if not d["has_text"]:
            show('<div class="note">No resume was uploaded, so Clarity and Edge are skipped and Match and Impact carry the score.</div>')
        total_w = sum(t["reqs"].values()) or 1
        srows = []
        for s, wt in sorted(t["reqs"].items(), key=lambda kv: -kv[1]):
            have = s in t["matched"]
            ev = t["per"].get(s)
            srows.append([esc(pretty(s)), f"{wt:.2f}", '<span class="chip ok">Have</span>' if have else '<span class="chip gap">Gap</span>',
                          f'{d["strengths"].get(s, 0):.0f}' if have else "—", "—" if ev is None else f"{ev:.0f}",
                          (f'+{wt / total_w * 100:.1f}' if have else f'{wt / total_w * 100:.1f} available')])
        show('<div class="sec">Skill-by-skill contribution to the Match score</div>')
        show(table(["Skill", "Weight", "Status", "Strength", "Semantic evidence", "Match points"], srows))
        gain = sorted(t["missing"], key=lambda s: -t["reqs"][s])[:3]
        if gain:
            show('<div class="note"><b>Biggest ways to raise your score:</b> learn ' + ", ".join(esc(pretty(s)) for s in gain) + ". Each adds Match points shown above.</div>")
        with st.expander("How the score is calculated"):
            st.markdown(
                "- **Match** (40%): weighted share of required skills you have. Earlier skills in a role, or must-haves in a job description, weigh more.\n"
                "- **Impact** (25%): average strength of your matched skills (frequency, concrete use, selection).\n"
                "- **Clarity** (15%): resume sections, email present, measurable results, and length.\n"
                "- **Edge** (20%): 70% semantic similarity between resume sentences and each required skill (SBERT, or TF-IDF fallback) plus 30% extra skills beyond the requirement.\n"
                "- Without a resume, Match counts 70% and Impact 30%.\n"
                "- Experience level does not change the score. It shapes the recommendations."
            )

    # ---------- Roadmap & learning ----------
    with tabs[2]:
        show('<div class="sec">Six-week career roadmap</div>')
        if d["roadmap"]:
            rrows = [[esc(r["when"]), esc(r["focus"]), esc(pretty(r["skill"])), esc(r["learn"]), esc(r["milestone"])] for r in d["roadmap"]]
            show(table(["When", "Focus", "Skill", "What to learn", "Milestone"], rrows))
            show('<div class="note">Ordered by the prioritization engine. Adjust the pace to your schedule.</div>')
            show('<div class="sec">Learning links</div>')
            lrows = []
            for p in d["priorities"][:9]:
                ln = learn_links(p["skill"])
                lrows.append([esc(pretty(p["skill"])), link(ln["docs"], "Official docs") if ln["docs"] else "—",
                              link(ln["course"], "Find courses"), link(ln["video"], "Video tutorials")])
            show(table(["Skill", "Docs", "Courses (Coursera)", "Videos (YouTube)"], lrows))
            show('<div class="note">Links open searches or official documentation, so you can pick a course that suits your budget and level.</div>')
        else:
            st.success("No gaps found. Focus on portfolio depth and interview practice.")

    # ---------- Recommendations ----------
    with tabs[3]:
        rc = d["recs"]
        show('<div class="sec">Personalized recommendations</div>')
        show('<div class="guide">' + "".join(f"<div>• {esc(x)}</div>" for x in rc["personal"]) + "</div>")
        cc1, cc2 = st.columns(2)
        with cc1:
            show('<div class="sec">Course recommendations</div>')
            show(table(["Skill", "Course focus"], [[esc(pretty(s)), esc(c)] for s, c in rc["courses"]]) if rc["courses"] else '<div class="note">No courses needed.</div>')
            show('<div class="sec">Certification recommendations</div>')
            show('<div class="guide">' + ("".join(f"<div>• {esc(x)}</div>" for x in rc["certs"]) or '<div class="note">No certification stands out for this target.</div>') + "</div>")
        with cc2:
            show('<div class="sec">Project recommendations</div>')
            show('<div class="guide">' + ("".join(f"<div>• {esc(x)}</div>" for x in rc["projects"]) or '<div class="note">Build a small project for your top missing skill.</div>') + "</div>")

    # ---------- Jobs: role matching + multi-job comparison ----------
    with tabs[4]:
        ranked = sorted(d["comparison"], key=lambda c: -c["overall"])
        names = [c["name"] for c in ranked]
        show('<div class="sec">Job-role matching</div>')
        show(table(["Job", "Overall", "Skill match", "Skills", "Top gaps", "Weeks"],
                   [[("★ " if c["name"] == t["name"] else "") + esc(c["name"]), bar(c["overall"]), f'{c["scores"]["match"]:.0f}%',
                     f'{len(c["matched"])}/{len(c["reqs"])}', esc(", ".join(pretty(s) for s in c["missing"][:3]) or "None"), str(c["weeks"])] for c in ranked]))
        show('<div class="note">★ marks your selected target. Roles use built-in skill lists; pasted job descriptions use skills found in their text.</div>')
        show('<div class="sec">Multi-job comparison</div>')
        chosen = st.multiselect("Choose jobs to compare", names, default=names[:3], key="cmp_jobs")
        if chosen:
            by = {c["name"]: c for c in ranked}
            head = "".join(f"<th>{esc(n)}</th>" for n in chosen)
            def row(label, fn):
                return f"<tr><td><b>{esc(label)}</b></td>" + "".join(f"<td>{fn(by[n])}</td>" for n in chosen) + "</tr>"
            body = (
                row("Overall", lambda c: bar(c["overall"]))
                + row("Skill match", lambda c: f'{c["scores"]["match"]:.0f}%')
                + row("Skill strength", lambda c: f'{c["scores"]["impact"]:.0f}%')
                + row("Skills matched", lambda c: f'{len(c["matched"])} of {len(c["reqs"])}')
                + row("Gaps", lambda c: chips(c["missing"][:4], "gap") or "None")
                + row("Weeks to close", lambda c: str(c["weeks"]))
            )
            show(f'<div class="tbl-wrap"><table class="tbl"><thead><tr><th>Measure</th>{head}</tr></thead><tbody>{body}</tbody></table></div>')

    # ---------- Market trends ----------
    with tabs[5]:
        show('<div class="sec">Role demand</div>')
        mrows = []
        for c in sorted(d["comparison"], key=lambda c: -c["overall"]):
            if c["name"] in ROLE_MARKET:
                dem, outlook, note = ROLE_MARKET[c["name"]]
                kind = "ok" if outlook == "Rising" else "neutral"
                mrows.append([esc(c["name"]), bar(dem, "b"), f'<span class="chip {kind}">{outlook}</span>', esc(note)])
        show(table(["Role", "Demand index", "Outlook", "Notes"], mrows) if mrows else '<div class="note">Select a built-in role to see role-level trends.</div>')
        show('<div class="sec">Skill demand for your target</div>')
        srows = []
        for s in sorted(t["reqs"], key=lambda s: -DEMAND.get(s, 55)):
            trend = "Rising" if s in RISING else "Cooling" if s in COOLING else "Stable"
            kind = "ok" if trend == "Rising" else "warn" if trend == "Cooling" else "neutral"
            have = '<span class="chip ok">Have</span>' if s in t["matched"] else '<span class="chip gap">Gap</span>'
            srows.append([esc(pretty(s)), bar(DEMAND.get(s, 55), "b"), f'<span class="chip {kind}">{trend}</span>', have])
        show(table(["Skill", "Demand index", "Trend", "You"], srows))
        show('<div class="note">Demand index, outlook, and trend labels are curated, indicative estimates built into this app. They are not live job-market data, so check current listings before making decisions.</div>')

    # ---------- Career switch simulator ----------
    with tabs[6]:
        show('<div class="gsub">Pick a role to switch to and the skills you plan to learn. See how your fit changes across every job.</div>')
        by = {c["name"]: c for c in d["comparison"]}
        sim_role = st.selectbox("Role to switch to", list(by.keys()), index=list(by.keys()).index(t["name"]), key="sim_role")
        sm = by[sim_role]
        show('<div class="sec">Transferable skills</div>')
        show(chips(sm["matched"], "ok") if sm["matched"] else '<div class="note">No direct overlap yet.</div>')
        all_missing = sorted({s for c in d["comparison"] for s in c["missing"]})
        default_learn = [s for s in sm["missing"][:2] if s in all_missing]
        learn = st.multiselect("Skills I plan to learn", [pretty(s) for s in all_missing],
                               default=[pretty(s) for s in default_learn], key="sim_learn")
        learn_c = {canon(x) for x in learn}
        new_skills = d["user_skills"] | learn_c
        new_strengths = dict(d["strengths"])
        for s in learn_c:
            new_strengths.setdefault(s, 50)
        clarity = d["parsed"]["clarity"] if d["parsed"] else None
        srows = []
        for c in sorted(d["comparison"], key=lambda c: -c["overall"]):
            nm = combine(c["reqs"], new_skills, new_strengths, clarity, c["semscore"])
            delta = nm["overall"] - c["overall"]
            srows.append([("★ " if c["name"] == sim_role else "") + esc(c["name"]), f'{c["overall"]:.0f}', f'{nm["overall"]:.0f}',
                          f'<span class="chip {"ok" if delta > 0.05 else "neutral"}">{delta:+.1f}</span>',
                          str(len(nm["missing"])), str(est_weeks(nm["missing"]))])
        show(table(["Job", "Now", "After learning", "Change", "Gaps left", "Weeks left"], srows))
        sel = combine(sm["reqs"], new_skills, new_strengths, clarity, sm["semscore"])
        a, b = tier(sel["overall"])
        show(f'<div class="glass"><div class="ghead">If you switch to {esc(sim_role)}</div>'
             f'<div class="guide">Projected score <b>{sel["overall"]:.0f}</b> ({a} / {b}). '
             f'You would still need {len(sel["missing"])} skill(s)'
             + (": " + esc(", ".join(pretty(s) for s in sel["missing"][:5])) if sel["missing"] else "") + ".</div>"
             '<div class="note">New skills are assumed to start at a basic level (strength 50). Resume-based parts of the score stay the same.</div></div>')

    # ---------- Readiness ----------
    with tabs[7]:
        r = d["readiness"]
        show(f'<div class="glass"><div class="ghead">Career readiness: {esc(r["level"])}</div>'
             f'<div class="guide">Estimated time to close your gaps: about <b>{r["weeks"]} weeks</b> at 8 to 10 hours a week (rough estimate).</div></div>')
        show('<div class="sec">Readiness checklist</div>')
        show("".join(f'<div class="ck"><span class="{"y" if ok else "n"}">{"✓" if ok else "✗"}</span><span>{esc(lbl)}</span></div>' for lbl, ok in r["checks"]))
        if r["core_missing"]:
            show('<div class="sec">Core skills still missing</div>')
            show(chips(r["core_missing"], "gap"))
        if r["note"]:
            show(f'<div class="note">{esc(r["note"])}</div>')
        show('<div class="sec">Career guidance</div>')
        show('<div class="guide">' + "".join(f"<div>{esc(p)}</div>" for p in d["guidance"]) + "</div>")

    # ---------- Resume parsing & JD analysis ----------
    with tabs[8]:
        show('<div class="sec">Resume parsing</div>')
        p = d["parsed"]
        if p:
            show("".join([
                f'<div class="kpis"><div class="kpi"><div class="kpi-l">Clarity</div><div class="kpi-v">{p["clarity"]:.0f}</div></div>',
                f'<div class="kpi"><div class="kpi-l">Words</div><div class="kpi-v">{p["words"]}</div></div>',
                f'<div class="kpi"><div class="kpi-l">Measurable results</div><div class="kpi-v">{p["quantified"]}</div></div>',
                f'<div class="kpi"><div class="kpi-l">Skills detected</div><div class="kpi-v">{len(d["user_skills"])}</div></div></div>',
            ]))
            show('<div class="note">Sections found</div>' + (chips_plain := "".join(f'<span class="chip ok">{esc(s.title())}</span>' for s in p["sections"]) or '<span class="note">None detected</span>'))
            missing_sec = [s for s in KEY_SECTIONS if s not in p["sections"]]
            if missing_sec:
                show('<div class="note">Sections not found</div>' + "".join(f'<span class="chip gap">{esc(s.title())}</span>' for s in missing_sec))
            extra_bits = [f'Email: {"found" if p["has_email"] else "not found"}']
            if p["years"]:
                extra_bits.append(f'Years mentioned: {p["years"][0]} to {p["years"][-1]}')
            show('<div class="note">' + esc(" · ".join(extra_bits)) + "</div>")
            show('<div class="sec">Skills detected from your resume and selections</div>')
            show(chips(sorted(d["user_skills"]), "neutral"))
        else:
            show('<div class="note">No resume uploaded. Upload a PDF, DOCX or TXT to see parsing results.</div>')

        show('<div class="sec">Job description analysis</div>')
        j = d["jd"]
        if j["text"] and (j["must"] or j["nice"]):
            show('<div class="note">Must-have skills</div>' + chips(sorted(j["must"]), "ok"))
            if j["nice"]:
                show('<div class="note">Nice-to-have skills</div>' + chips(sorted(j["nice"]), "info"))
            jm = next((c for c in d["comparison"] if c["name"] in ("Custom job description", "Pasted job description")), None)
            if jm:
                show(f'<div class="note">Your fit for this job description: <b>{jm["overall"]:.0f}</b> overall, with {len(jm["matched"])} of {len(jm["reqs"])} skills matched.</div>')
            if j["years"]:
                show(f'<div class="note">Experience asked: about {j["years"]}+ years. Your level: {esc(d["experience"])}.</div>')
        else:
            show('<div class="note">Paste a job description above to see its must-have and nice-to-have skills.</div>')

    # ---------- Download ----------
    with tabs[9]:
        st.write("Save your results.")
        st.download_button("Download PDF report", data=build_pdf(d), file_name="skillbridge_report.pdf", mime="application/pdf")
        prio = {p["skill"]: p["priority"] for p in d["priorities"]}
        csv = pd.DataFrame({
            "Skill": [pretty(s) for s in t["reqs"]],
            "Weight": [t["reqs"][s] for s in t["reqs"]],
            "Status": ["Have" if s in t["matched"] else "Gap" for s in t["reqs"]],
            "Strength": [round(d["strengths"].get(s, 0)) if s in t["matched"] else "" for s in t["reqs"]],
            "Priority": [prio.get(s, "") for s in t["reqs"]],
        }).to_csv(index=False).encode("utf-8")
        st.download_button("Download skill checklist (CSV)", data=csv, file_name="skillbridge_skills.csv", mime="text/csv")

    st.caption("Skills are found by keyword matching and semantic similarity, so synonyms and skills outside the built-in list can be missed. "
               "Scores are rough indicators, not hiring predictions. Your resume is only processed during this session.")
